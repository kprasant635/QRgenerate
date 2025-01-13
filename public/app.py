import qrcode #Import The Module

# import module 
import openpyxl 


def increase_quota():
	file1 = open("quota.txt", "r+")
	val=file1.read()
	print(val)
	new_val=int(val)+1
	print(new_val)
	file1.truncate()
	file1.write(str(new_val))
	file1.close()

def qr_generator(FNAME,LNAME,TEL_HOME,TEL_WORK,EMAIL,ORG,TITLE,ADDRESS,URL):   

	vcard='''BEGIN:VCARD
N:'''+str(FNAME)+''';'''+str(LNAME)+''';
TEL;TYPE=work,VOICE:'''+str(TEL_WORK)+'''
TEL;TYPE=home,VOICE:'''+str(TEL_HOME)+'''
TEL;TYPE=fax:
EMAIL:'''+EMAIL+'''
ORG:'''+str(ORG)+'''
TITLE:'''+str(TITLE)+'''
ADR;TYPE=WORK,PREF:;;'''+str(ADDRESS)+'''
URL:'''+str(URL)+'''
VERSION:3.0
END:VCARD''' 

	img = qrcode.make(vcard)

	type(img)
	filename=str(TEL_WORK)+"_"+FNAME+".png"
	img.save(filename) #Save the file
	# increase_quota()

# load excel with its path 
wrkbk = openpyxl.load_workbook("sample_qr.xlsx") 
  
sh = wrkbk.active 
  
# iterate through excel and display data 
for i in range(2, sh.max_row+1): 
    print("\n") 
    print("Row ", i, " data :") 
      
    for j in range(1, sh.max_column+1): 
        cell_obj = sh.cell(row=i, column=j) 
        print(cell_obj.value, end=" ")
       
        if j==1:
        	FNAME=cell_obj.value
        if j==2:
        	LNAME=cell_obj.value
        if j==3:
        	TEL_HOME=cell_obj.value
        if j==4:
        	TEL_WORK=cell_obj.value
        if j==5:
        	EMAIL=cell_obj.value
        if j==6:
        	ORG=cell_obj.value
        if j==7:
        	TITLE=cell_obj.value
        if j==8:
        	ADDRESS=cell_obj.value
        if j==9:
        	URL=cell_obj.value
        	qr_generator(FNAME,LNAME,TEL_HOME,TEL_WORK,EMAIL,ORG,TITLE,ADDRESS,URL)

